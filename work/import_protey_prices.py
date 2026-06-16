from __future__ import annotations

import datetime as dt
import hashlib
import json
import sqlite3
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import openpyxl


SOURCE = Path(r"C:\Users\admla\Desktop\Протей\Протей.xlsx")
DB_PATH = Path(r"C:\Users\admla\Documents\Codex\2026-06-02\files-mentioned-by-the-user-02\outputs\cvsig_prices_2026-06-02.sqlite")
SYSTEM = "Протей"
PRICE_DATE = "2026-06-02"
VAT_NOTE = "по исходной таблице: руб. с НДС"


TOTAL_PREFIXES = (
    "прочее, итого",
    "итого",
    "всего",
    "сумма ндс",
)


def money_to_kopeks(value):
    if value is None or value == "":
        return None
    amount = Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    return int(amount * 100)


def num_or_none(value):
    if value is None or value == "":
        return None
    return float(value)


def text_or_none(value):
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def is_total_row(name: str) -> bool:
    normalized = name.strip().lower()
    return normalized.startswith(TOTAL_PREFIXES)


def iter_rows():
    workbook = openpyxl.load_workbook(SOURCE, data_only=True, read_only=True)
    for sheet in workbook.worksheets:
        for excel_row, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if excel_row <= 4:
                continue

            position = row[0] if len(row) > 0 else None
            name = row[1] if len(row) > 1 else None
            if not isinstance(position, (int, float)) or not isinstance(name, str):
                continue
            if is_total_row(name):
                continue

            material_unit = row[6] if len(row) > 6 else None
            material_total = row[7] if len(row) > 7 else None
            installation_unit = row[8] if len(row) > 8 else None
            installation_total = row[9] if len(row) > 9 else None
            has_price = any(
                value is not None
                for value in (material_unit, material_total, installation_unit, installation_total)
            )
            if not has_price:
                continue

            yield {
                "sheet_name": sheet.title,
                "excel_row": excel_row,
                "position": int(position) if float(position).is_integer() else position,
                "name": text_or_none(name),
                "type_code_article": text_or_none(row[2] if len(row) > 2 else None),
                "manufacturer": text_or_none(row[3] if len(row) > 3 else None),
                "unit": text_or_none(row[4] if len(row) > 4 else None),
                "quantity": num_or_none(row[5] if len(row) > 5 else None),
                "material_unit_kopeks": money_to_kopeks(material_unit),
                "material_total_kopeks": money_to_kopeks(material_total),
                "installation_unit_kopeks": money_to_kopeks(installation_unit),
                "installation_total_kopeks": money_to_kopeks(installation_total),
            }


def main():
    rows = list(iter_rows())
    source_stat = SOURCE.stat()
    source_hash = file_sha256(SOURCE)
    headers = {
        "columns": [
            "поз.",
            "Наименование",
            "Тип / код / артикул",
            "Изготовитель",
            "Ед.",
            "Кол-во",
            "Стоимость материала: ед. расценка, руб. с НДС",
            "Стоимость материала: итого, руб. с НДС",
            "Стоимость монтажа: ед. расценка, руб. с НДС",
            "Стоимость монтажа: итого, руб. с НДС",
        ],
        "system_assigned": SYSTEM,
        "skipped_summary_rows": list(TOTAL_PREFIXES),
    }

    with sqlite3.connect(DB_PATH) as conn:
        conn.execute("PRAGMA foreign_keys = ON")
        previous = conn.execute(
            "SELECT id FROM price_snapshots WHERE source_file = ? OR source_sha256 = ?",
            (str(SOURCE), source_hash),
        ).fetchall()
        for (snapshot_id,) in previous:
            conn.execute("DELETE FROM price_snapshots WHERE id = ?", (snapshot_id,))

        cursor = conn.execute(
            """
            INSERT INTO price_snapshots (
                source_file,
                source_sha256,
                source_last_modified,
                price_date,
                imported_at,
                source_headers_json
            )
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (
                str(SOURCE),
                source_hash,
                dt.datetime.fromtimestamp(source_stat.st_mtime).isoformat(timespec="seconds"),
                PRICE_DATE,
                dt.datetime.now().isoformat(timespec="seconds"),
                json.dumps(headers, ensure_ascii=False),
            ),
        )
        snapshot_id = cursor.lastrowid

        for row in rows:
            cursor = conn.execute(
                """
                INSERT INTO price_lines (
                    snapshot_id,
                    system,
                    excel_row,
                    position,
                    name,
                    type_code_article,
                    manufacturer,
                    unit,
                    quantity,
                    note
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    snapshot_id,
                    SYSTEM,
                    row["excel_row"],
                    str(row["position"]),
                    row["name"],
                    row["type_code_article"],
                    row["manufacturer"],
                    row["unit"],
                    row["quantity"],
                    row["sheet_name"],
                ),
            )
            line_id = cursor.lastrowid

            if row["material_unit_kopeks"] is not None or row["material_total_kopeks"] is not None:
                conn.execute(
                    """
                    INSERT INTO prices (
                        line_id,
                        price_kind,
                        unit_price_kopeks,
                        total_price_kopeks,
                        vat_note
                    )
                    VALUES (?, 'material', ?, ?, ?)
                    """,
                    (
                        line_id,
                        row["material_unit_kopeks"],
                        row["material_total_kopeks"],
                        VAT_NOTE,
                    ),
                )

            if row["installation_unit_kopeks"] is not None or row["installation_total_kopeks"] is not None:
                conn.execute(
                    """
                    INSERT INTO prices (
                        line_id,
                        price_kind,
                        unit_price_kopeks,
                        total_price_kopeks,
                        vat_note
                    )
                    VALUES (?, 'installation', ?, ?, ?)
                    """,
                    (
                        line_id,
                        row["installation_unit_kopeks"],
                        row["installation_total_kopeks"],
                        VAT_NOTE,
                    ),
                )

        result = {
            "inserted_lines": conn.execute(
                "SELECT COUNT(*) FROM price_lines WHERE snapshot_id = ?", (snapshot_id,)
            ).fetchone()[0],
            "inserted_prices": conn.execute(
                """
                SELECT price_kind, COUNT(*), ROUND(SUM(total_price_kopeks) / 100.0, 2)
                FROM prices
                JOIN price_lines ON price_lines.id = prices.line_id
                WHERE snapshot_id = ?
                GROUP BY price_kind
                ORDER BY price_kind
                """,
                (snapshot_id,),
            ).fetchall(),
            "deleted_previous_snapshots": len(previous),
        }

    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
