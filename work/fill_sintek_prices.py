from __future__ import annotations

import json
import shutil
from copy import copy
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


WORKBOOK_PATH = Path(r"C:\Users\admla\Desktop\Синтек 0806.xlsx")
REPORT_JSON = Path(r"C:\Users\admla\Documents\Codex\2026-06-02\files-mentioned-by-the-user-02\work\sintek_match_report.json")
BACKUP_PATH = WORKBOOK_PATH.with_name("Синтек 0806.backup-before-prices.xlsx")
OUTPUT_PATH = Path(r"C:\Users\admla\Documents\Codex\2026-06-02\files-mentioned-by-the-user-02\outputs\Синтек 0806 с ценами.xlsx")
OUTPUT_VAT_PATH = Path(r"C:\Users\admla\Documents\Codex\2026-06-02\files-mentioned-by-the-user-02\outputs\Синтек 0806 с ценами с НДС.xlsx")
REPORT_SHEET = "Отчет сопоставления"


def round_money(value):
    if value is None:
        return None
    return round(float(value), 2)


def copy_format(source, targets):
    for target in targets:
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.alignment:
            target.alignment = copy(source.alignment)


def write_report_sheet(workbook, report):
    if REPORT_SHEET in workbook.sheetnames:
        del workbook[REPORT_SHEET]
    sheet = workbook.create_sheet(REPORT_SHEET)

    headers = [
        "Статус",
        "Лист",
        "Строка",
        "Поз.",
        "Наименование в документе",
        "Тип/код в документе",
        "Кол-во",
        "Материал ед., с НДС",
        "Монтаж ед., с НДС",
        "Источник",
        "Совпавшее наименование",
        "Совпавший тип/код",
        "Балл",
        "Отступ",
        "НДС источника",
        "Комментарий",
    ]
    sheet.append(headers)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for row in report:
        filled = row["action"] == "fill"
        vat_notes = sorted(
            {
                note
                for note in (row.get("material_vat_note"), row.get("installation_vat_note"))
                if note
            }
        )
        comment = "Заполнено автоматически" if filled else "Не заполнено: требуется ручная проверка"

        sheet.append(
            [
                "заполнено" if filled else "проверить",
                row["sheet"],
                row["row"],
                row["position"],
                row["target_name"],
                row["target_type"],
                row["target_qty"],
                row["material_unit_with_vat"],
                row["installation_unit_with_vat"],
                row["match_system"],
                row["match_name"],
                row["match_type"],
                row["score"],
                row["margin"],
                "; ".join(vat_notes),
                comment,
            ]
        )

    for row in sheet.iter_rows(min_row=2):
        status = row[0].value
        fill = PatternFill("solid", fgColor="E2F0D9" if status == "заполнено" else "FFF2CC")
        for cell in row:
            cell.fill = fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    widths = {
        1: 12,
        2: 18,
        3: 8,
        4: 8,
        5: 45,
        6: 28,
        7: 10,
        8: 16,
        9: 16,
        10: 12,
        11: 38,
        12: 28,
        13: 10,
        14: 10,
        15: 28,
        16: 42,
    }
    for col_idx, width in widths.items():
        sheet.column_dimensions[get_column_letter(col_idx)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def main():
    report = json.loads(REPORT_JSON.read_text(encoding="utf-8"))
    if not BACKUP_PATH.exists():
        shutil.copy2(WORKBOOK_PATH, BACKUP_PATH)

    workbook = openpyxl.load_workbook(WORKBOOK_PATH)
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        if sheet.max_column >= 10 and sheet.max_row >= 1:
            for col in (7, 8, 9, 10):
                value = sheet.cell(1, col).value
                if isinstance(value, str):
                    sheet.cell(1, col).value = value.replace("без НДС", "с НДС")

    filled_rows = 0
    filled_cells = 0

    for row in report:
        if row["action"] != "fill":
            continue

        sheet = workbook[row["sheet"]]
        excel_row = int(row["row"])
        quantity = row["target_qty"]
        material_unit = round_money(row["material_unit_with_vat"])
        installation_unit = round_money(row["installation_unit_with_vat"])

        # Keep the row's existing style pattern in newly populated price cells.
        copy_format(sheet.cell(excel_row, 6), [sheet.cell(excel_row, col) for col in range(7, 11)])

        row_filled = False
        if material_unit is not None:
            sheet.cell(excel_row, 7).value = material_unit
            sheet.cell(excel_row, 8).value = round_money(material_unit * float(quantity)) if quantity else None
            row_filled = True
            filled_cells += 2

        if installation_unit is not None:
            sheet.cell(excel_row, 9).value = installation_unit
            sheet.cell(excel_row, 10).value = round_money(installation_unit * float(quantity)) if quantity else None
            row_filled = True
            filled_cells += 2

        for col in range(7, 11):
            sheet.cell(excel_row, col).number_format = '#,##0.00'

        source_note = (
            f"Цена из базы: {row['match_system']} | {row['match_name']} | {row['match_type'] or ''}. "
            f"Обновлено {datetime.now().strftime('%Y-%m-%d %H:%M')}."
        )
        sheet.cell(excel_row, 11).value = source_note

        if row_filled:
            filled_rows += 1

    write_report_sheet(workbook, report)
    OUTPUT_VAT_PATH.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT_VAT_PATH)
    saved_path = OUTPUT_VAT_PATH

    result = {
        "workbook": str(saved_path),
        "source_workbook": str(WORKBOOK_PATH),
        "backup": str(BACKUP_PATH),
        "filled_rows": filled_rows,
        "filled_cells": filled_cells,
        "report_sheet": REPORT_SHEET,
        "review_rows": sum(1 for row in report if row["action"] != "fill"),
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
