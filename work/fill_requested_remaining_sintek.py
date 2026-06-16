from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]


def is_blank(value):
    return value is None or value == ""


def as_number(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return value
    try:
        return float(str(value).replace(" ", "").replace(",", "."))
    except ValueError:
        return None


def set_money(cell, value):
    if value is None:
        return
    if abs(value - round(value)) < 0.000001:
        cell.value = int(round(value))
    else:
        cell.value = round(value, 2)


def append_note(cell, note):
    if not note:
        return
    if is_blank(cell.value):
        cell.value = note
        return
    text = str(cell.value)
    if note not in text:
        cell.value = f"{text} | {note}"


def fill_row(ws, row, material_unit=None, install_unit=None, source_note=""):
    qty = as_number(ws.cell(row, 6).value)
    g = ws.cell(row, 7)
    h = ws.cell(row, 8)
    i = ws.cell(row, 9)
    j = ws.cell(row, 10)
    k = ws.cell(row, 11)

    changed = []

    if is_blank(g.value) and material_unit is not None:
        set_money(g, material_unit)
        changed.append("материал за ед.")

    current_material = as_number(g.value)
    if is_blank(h.value) and current_material is not None and qty is not None:
        set_money(h, current_material * qty)
        changed.append("материал сумма")

    if is_blank(i.value) and install_unit is not None:
        set_money(i, install_unit)
        changed.append("монтаж за ед.")

    current_install = as_number(i.value)
    if is_blank(j.value) and current_install is not None and qty is not None:
        set_money(j, current_install * qty)
        changed.append("монтаж сумма")

    if changed:
        append_note(k, source_note)

    return changed


def sync_report(wb, filled):
    if len(wb.worksheets) < 2:
        return
    report = wb.worksheets[-1]
    headers = {}
    for col in range(1, report.max_column + 1):
        value = report.cell(1, col).value
        if value:
            headers[str(value).strip().lower()] = col

    row_col = headers.get("строка")
    status_col = headers.get("статус")
    mat_col = headers.get("материал, руб./ед. с ндс")
    inst_col = headers.get("монтаж, руб./ед. с ндс")
    comment_col = headers.get("комментарий")
    sheet_col = headers.get("лист")

    if not row_col or not status_col:
        return

    by_key = {(ws_title, row): data for ws_title, row, data in filled}
    for r in range(2, report.max_row + 1):
        row_num = report.cell(r, row_col).value
        try:
            row_num = int(row_num)
        except (TypeError, ValueError):
            continue
        sheet_name = report.cell(r, sheet_col).value if sheet_col else None
        data = by_key.get((sheet_name, row_num))
        if not data:
            continue
        report.cell(r, status_col).value = "дозаполнено по запросу"
        if mat_col:
            report.cell(r, mat_col).value = data["material"]
        if inst_col:
            report.cell(r, inst_col).value = data["install"]
        if comment_col:
            append_note(report.cell(r, comment_col), data["note"])


def main():
    files = sorted((ROOT / "outputs").glob("*0806*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True)
    if not files:
        raise FileNotFoundError("Не найден файл Синтек 0806 в папке outputs")

    path = files[0]
    wb = load_workbook(path)
    ws = wb.worksheets[0]

    rows = {
        20: {
            "material_unit": None,
            "install_unit": 2000,
            "note": "Дозаполнено: Дагестан / АПС; оповещатель комбинированный свето-звуковой; монтаж 2000 руб. с НДС",
        },
        21: {
            "material_unit": None,
            "install_unit": 750,
            "note": "Дозаполнено: база 2026-04-24; оповещатель охранно-пожарный звуковой; монтаж 750 руб. с НДС",
        },
        29: {
            "material_unit": None,
            "install_unit": 3000,
            "note": "Дозаполнено: база 2026-04-24; комплект ОКЛ СМО d19-20; монтаж 3000 руб. с НДС",
        },
        33: {
            "material_unit": 51.82,
            "install_unit": 93,
            "note": "Дозаполнено: база 2026-04-24; струбцина монтажная М8; материал 51.82, монтаж 93 руб. с НДС",
        },
        36: {
            "material_unit": 48.32,
            "install_unit": 10,
            "note": "Дозаполнено: Дагестан; гайка соединительная М8; материал 48.32, монтаж 10 руб. с НДС",
        },
        37: {
            "material_unit": 1.30,
            "install_unit": 10,
            "note": "Дозаполнено: база 2026-04-24; гайка DIN 934 М8; материал 1.30, монтаж 10 руб. с НДС",
        },
        38: {
            "material_unit": 7,
            "install_unit": 10,
            "note": "Дозаполнено: РВС; гайка М6 с насечкой как ближайший аналог самоконтрящейся М6; материал 7, монтаж 10 руб. с НДС",
        },
        73: {
            "material_unit": None,
            "install_unit": 3000,
            "note": "Дозаполнено: база 2026-04-24; комплект ОКЛ; монтаж 3000 руб. с НДС",
        },
        81: {
            "material_unit": 5.50,
            "install_unit": 10,
            "note": "Дозаполнено: Дагестан; гайка М8 как ближайший аналог самоконтрящейся М8; материал 5.50, монтаж 10 руб. с НДС",
        },
    }

    filled = []
    for row, data in rows.items():
        changed = fill_row(
            ws,
            row,
            material_unit=data["material_unit"],
            install_unit=data["install_unit"],
            source_note=data["note"],
        )
        if changed:
            filled.append(
                (
                    ws.title,
                    row,
                    {
                        "material": ws.cell(row, 7).value,
                        "install": ws.cell(row, 9).value,
                        "note": data["note"],
                        "changed": changed,
                    },
                )
            )

    sync_report(wb, filled)
    wb.save(path)

    print(path)
    print(filled)


if __name__ == "__main__":
    main()
