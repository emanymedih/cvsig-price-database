"""Turn a specification (спецификация: Позиция / Наименование / Тип / Код / Завод / Ед. / Кол-во ...)
into a КП workbook in the house template (same layout as the СПС/СОУЭ КП of 18.10.2026):

    python3 work/spec_to_kp.py SPEC.xlsx -o KP.xlsx --title "Перечень ..." --object "Объект: ..."

Sections (rows with text only in the first column) are kept; positions are renumbered "section.item".
Price columns stay empty: fill them with work/fill_estimate.py. Rates of the "прочие расходы"
(3/2/6/2/5 %) and VAT 22 % are the template defaults and can be changed with --rates / --vat.
"""
from __future__ import annotations

import argparse
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Times New Roman"
HEAD = PatternFill("solid", fgColor="FCE8DA")
THIN = Side(style="thin", color="000000")
GRID = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY = '#,##0.00 "₽"'  # fill_estimate.py restyles it into the accounting format
WIDTHS = {"A": 8, "B": 74, "C": 43, "D": 24, "E": 9, "F": 11, "G": 16, "H": 17, "I": 16, "J": 17, "K": 29}
EXPENSES = [  # (name, rate, base)
    ("Транспортные расходы", 0.03, "Материалы"),
    ("Расходные материалы", 0.02, "Материалы"),
    ("Пусконаладочные работы", 0.06, "Монтажные работы"),
    ("Непредвиденные расходы", 0.02, "Полная стоимость"),
    ("Комплект исполнительной документации", 0.05, "Монтажные работы"),
]


def read_spec(path: Path):
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    head = None
    for r in range(1, 15):
        vals = [str(ws.cell(r, c).value or "").strip().lower() for c in range(1, ws.max_column + 1)]
        if any(v.startswith("наименование") for v in vals):
            head = r
            cols = {}
            for c, v in enumerate(vals, 1):
                if v.startswith("поз"):
                    cols["pos"] = c
                elif v.startswith("наименование"):
                    cols["name"] = c
                elif v.startswith("тип") or "марка" in v:
                    cols["type"] = c
                elif v.startswith("код"):
                    cols["code"] = c
                elif v.startswith("завод") or "изготовит" in v or "производит" in v:
                    cols["maker"] = c
                elif v.startswith("ед"):
                    cols["unit"] = c
                elif v.startswith("кол"):
                    cols["qty"] = c
                elif v.startswith("примеч"):
                    cols["note"] = c
            break
    if head is None:
        raise SystemExit("В спецификации не найдена строка заголовков")
    blocks, current = [], None
    for r in range(head + 1, ws.max_row + 1):
        get = lambda k: ws.cell(r, cols[k]).value if k in cols else None  # noqa: E731
        name, qty = get("name"), get("qty")
        first = ws.cell(r, 1).value
        if isinstance(first, str) and not name and qty is None:
            current = {"title": first.strip(), "items": []}
            blocks.append(current)
            continue
        if not name or not isinstance(qty, (int, float)) or str(name).strip().isdigit():
            continue  # empty rows and the "1 2 3 ..." column-number row of a ГОСТ specification
        if current is None:
            current = {"title": "", "items": []}
            blocks.append(current)
        typ, code = str(get("type") or "").strip(), str(get("code") or "").strip()
        current["items"].append({
            "name": str(name).strip(), "article": typ + (f"; код {code}" if code and code != typ else ""),
            "maker": str(get("maker") or "").strip(), "unit": str(get("unit") or "").strip(), "qty": qty,
            "note": str(get("note") or "").strip(), "spec_pos": get("pos"),
        })
    return blocks


def cell(ws, ref, value=None, bold=False, italic=False, fill=None, h="center", border=True, fmt=None, size=12):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = Font(name=FONT, size=size, bold=bold, italic=italic)
    c.alignment = Alignment(horizontal=h, vertical="center", wrap_text=True)
    if fill:
        c.fill = fill
    if border:
        c.border = GRID
    if fmt:
        c.number_format = fmt
    return c


def band(ws, r, text, bold=True, italic=True, fill=HEAD, h="center", last="K"):
    ws.merge_cells(f"A{r}:{last}{r}")
    for col in "ABCDEFGHIJK"[: "ABCDEFGHIJK".index(last) + 1]:
        cell(ws, f"{col}{r}", bold=bold, italic=italic, fill=fill)
    cell(ws, f"A{r}", text, bold=bold, italic=italic, fill=fill, h=h)


def build(blocks, title: str, obj: str, rates, vat: float) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "КП"
    for k, w in WIDTHS.items():
        ws.column_dimensions[k].width = w
    ws.merge_cells("A1:K2")
    cell(ws, "A1", title, bold=True, border=False, size=16)
    ws.merge_cells("A3:K3")
    cell(ws, "A3", obj, border=False)
    for r, h in ((1, 30), (2, 30), (3, 27), (4, 30), (5, 30)):
        ws.row_dimensions[r].height = h
    for col, text in zip("ABCDEF", ["Поз.", "Наименование и техническая характеристика", "Тип / марка / артикул",
                                     "Изготовитель", "Ед.", "Кол-во"]):
        ws.merge_cells(f"{col}4:{col}5")
        cell(ws, f"{col}4", text, bold=True, fill=HEAD)
        cell(ws, f"{col}5", bold=True, fill=HEAD)
    ws.merge_cells("K4:K5")
    cell(ws, "K4", "Примечание", bold=True, fill=HEAD)
    cell(ws, "K5", bold=True, fill=HEAD)
    for a, b, text in (("G", "H", "Стоимость материала"), ("I", "J", "Стоимость монтажа")):
        ws.merge_cells(f"{a}4:{b}4")
        cell(ws, f"{a}4", text, bold=True, fill=HEAD)
        cell(ws, f"{b}4", bold=True, fill=HEAD)
        cell(ws, f"{a}5", "за ед. с НДС, руб.", bold=True, fill=HEAD)
        cell(ws, f"{b}5", "итого с НДС, руб.", bold=True, fill=HEAD)

    r, first = 6, None
    for s, block in enumerate(blocks, 1):
        if block["title"]:
            band(ws, r, block["title"])
            ws.row_dimensions[r].height = 27
            r += 1
        for i, it in enumerate(block["items"], 1):
            vals = {"A": f"{s}.{i}", "B": it["name"], "C": it["article"] or None, "D": it["maker"] or None,
                    "E": it["unit"], "F": it["qty"], "K": it["note"] or None}
            for col in "ABCDEFGHIJK":
                cell(ws, f"{col}{r}", vals.get(col), h="left" if col == "B" else "center",
                     fmt="#,##0.###" if col == "F" else MONEY if col in "GHIJ" else None)
            ws[f"H{r}"] = f'=IF(ISNUMBER(G{r}),ROUND(F{r}*G{r},2),"")'
            ws[f"J{r}"] = f'=IF(ISNUMBER(I{r}),ROUND(F{r}*I{r},2),"")'
            first = first or r
            r += 1
    last = r - 1
    tot = r
    ws.merge_cells(f"A{tot}:F{tot}")
    for col in "ABCDEFGHIJK":
        cell(ws, f"{col}{tot}", bold=True, fill=HEAD, fmt=MONEY if col in "GHIJ" else None)
    cell(ws, f"A{tot}", "Итого по заполненным расценкам", bold=True, fill=HEAD, h="right")
    ws[f"H{tot}"] = f'=IF(COUNT(G{first}:G{last})=0,"",SUM(H{first}:H{last}))'
    ws[f"J{tot}"] = f'=IF(COUNT(I{first}:I{last})=0,"",SUM(J{first}:J{last}))'

    r = tot + 1
    band(ws, r, "Прочие расходы")
    r += 1
    hdr = r
    for col, text in zip("ABCDEF", ["№", "Наименование расходов", "Ставка", "База расчёта", "Ед.", "Кол-во"]):
        cell(ws, f"{col}{r}", text, bold=True, fill=HEAD)
    ws.merge_cells(f"G{r}:J{r}")
    for col in "GHIJ":
        cell(ws, f"{col}{r}", bold=True, fill=HEAD)
    cell(ws, f"K{r}", "Сумма с НДС, руб.", bold=True, fill=HEAD)
    rows = {}
    for n, ((name, _, base), rate) in enumerate(zip(EXPENSES, rates), 1):
        r += 1
        rows[n] = r
        for col, v in zip("ABCDEF", [f"{len(blocks) + 1}.{n}", name, rate, base, "к-т", 1]):
            cell(ws, f"{col}{r}", v, h="left" if col in "BD" else "center", fmt="0%" if col == "C" else None)
        for col in "GHIJ":
            cell(ws, f"{col}{r}")
        cell(ws, f"K{r}", h="right", fmt=MONEY)
    H, J = f"H{tot}", f"J{tot}"
    k = {n: f"K{rows[n]}" for n in rows}
    c = {n: f"C{rows[n]}" for n in rows}
    ws[k[1]] = f'=IF(ISNUMBER({H}),ROUND({H}*{c[1]},2),"")'
    ws[k[2]] = f'=IF(ISNUMBER({H}),ROUND({H}*{c[2]},2),"")'
    ws[k[3]] = f'=IF(ISNUMBER({J}),ROUND({J}*{c[3]},2),"")'
    ws[k[4]] = f'=IF(COUNT({H},{J})=0,"",ROUND(SUM({H},{J},{k[1]}:{k[3]},{k[5]})*{c[4]}/(1-{c[4]}),2))'
    ws[k[5]] = f'=IF(ISNUMBER({J}),ROUND({J}*{c[5]},2),"")'

    def total_row(rr, text, formula, fill=None):
        ws.merge_cells(f"A{rr}:J{rr}")
        for col in "ABCDEFGHIJ":
            cell(ws, f"{col}{rr}", bold=True, fill=fill, h="left")
        cell(ws, f"A{rr}", text, bold=True, fill=fill, h="left")
        cell(ws, f"K{rr}", formula, bold=True, fill=fill, h="right", fmt=MONEY)

    r += 1
    other = r
    total_row(r, "Прочие, итого:", f'=IF(COUNT(K{rows[1]}:K{rows[5]})=0,"",SUM(K{rows[1]}:K{rows[5]}))', HEAD)
    total_row(r + 1, "Итого с НДС, общая стоимость материалов:", f'=IF(ISNUMBER({H}),{H},"")')
    total_row(r + 2, "Итого с НДС, общая стоимость монтажных работ:", f'=IF(ISNUMBER({J}),{J},"")')
    total_row(r + 3, "Итого с НДС, прочие расходы:", f'=IF(ISNUMBER(K{other}),K{other},"")')
    total_row(r + 4, "Всего с НДС, руб.:", f'=IF(COUNT(K{r + 1}:K{r + 3})=0,"",SUM(K{r + 1}:K{r + 3}))', HEAD)
    v = r + 5
    ws.merge_cells(f"A{v}:B{v}")
    ws.merge_cells(f"D{v}:J{v}")
    for col in "ABCDEFGHIJ":
        cell(ws, f"{col}{v}", bold=True, h="left")
    cell(ws, f"A{v}", "Сумма НДС", bold=True, h="left")
    cell(ws, f"C{v}", vat, bold=True, fmt="0%")
    cell(ws, f"D{v}", "В составе общей стоимости", bold=True, h="left")
    cell(ws, f"K{v}", f'=IF(ISNUMBER(K{r + 4}),ROUND(K{r + 4}*C{v}/(1+C{v}),2),"")', bold=True, h="right", fmt=MONEY)
    for i, text in enumerate([f"Цены вводятся с НДС {vat:.0%}. НДС выделен из общей суммы. Итоги рассчитаны по заполненным расценкам.",
                              "Непредвиденные расходы: 2% от полной стоимости, включая прочие расходы. "
                              "Ставки расходов и НДС доступны для изменения."], 1):
        ws.merge_cells(f"A{v + i}:K{v + i}")
        cell(ws, f"A{v + i}", text, italic=True, border=False, h="left")
        ws.row_dimensions[v + i].height = 25
    for rr in range(tot, v + 1):
        ws.row_dimensions[rr].height = 30
    ws.print_area = f"A1:K{v + 2}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.25
    ws.print_title_rows = "4:5"
    return wb


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("spec", type=Path)
    ap.add_argument("-o", "--output", type=Path, required=True)
    ap.add_argument("--title", default="Перечень оборудования, кабельных изделий и материалов")
    ap.add_argument("--object", default="Объект: ")
    ap.add_argument("--rates", default="0.03,0.02,0.06,0.02,0.05")
    ap.add_argument("--vat", type=float, default=0.22)
    args = ap.parse_args()
    blocks = read_spec(args.spec)
    wb = build(blocks, args.title, args.object, [float(x) for x in args.rates.split(",")], args.vat)
    wb.save(args.output)
    print(f"Разделов {len(blocks)}, позиций {sum(len(b['items']) for b in blocks)} -> {args.output}")


if __name__ == "__main__":
    main()
