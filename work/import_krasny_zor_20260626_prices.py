from __future__ import annotations

import datetime as dt
import hashlib
import json
import sqlite3
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

import openpyxl


SOURCE = Path(r"C:\Users\admla\Downloads\krasny_zor (26.06).xlsx")
DB_PATH = Path(
    r"C:\Users\admla\Documents\Codex\2026-06-02\files-mentioned-by-the-user-02\outputs\cvsig_prices_2026-06-02.sqlite"
)
PRICE_DATE = "2026-06-26"
SYSTEM_MAIN = "Красный Зор / АПС и СОУЭ"
SYSTEM_MISC = "Красный Зор / Прочие расходы"
VAT_NOTE = "по КП: руб. с НДС"


def money_to_kopeks(value):
    if value is None or value == "":
        return None
    try:
        amount = Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError):
        return None
    return int(amount * 100)


def num_or_none(value):
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def text_or_none(value):
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def clean_position(value, excel_row):
    if value is None or value == "":
        return f"row-{excel_row}"
    if isinstance(value, bool):
        return f"row-{excel_row}"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value)
    return str(value).strip()


def derive_total(unit_kopeks, quantity):
    if unit_kopeks is None or quantity is None:
        return None
    amount = Decimal(unit_kopeks) * Decimal(str(quantity))
    return int(amount.quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def iter_price_rows(path: Path):
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheet = workbook["KP"]
    for excel_row, row in enumerate(
        sheet.iter_rows(min_row=10, max_row=sheet.max_row, max_col=11, values_only=True),
        start=10,
    ):
        name = text_or_none(row[1])
        if not name:
            continue

        position = clean_position(row[0], excel_row)
        unit = text_or_none(row[4])
        quantity = num_or_none(row[5])
        material_unit = money_to_kopeks(row[6])
        material_total = money_to_kopeks(row[7])
        installation_unit = money_to_kopeks(row[8])
        installation_total = money_to_kopeks(row[9])

        if excel_row <= 47:
            system = SYSTEM_MAIN
            if not unit or quantity is None:
                continue
        elif 50 <= excel_row <= 54:
            system = SYSTEM_MISC
            unit = unit or "к-т"
            quantity = quantity if quantity is not None else 1.0
            material_total = material_total if material_total is not None else derive_total(material_unit, quantity)
            installation_total = (
                installation_total
                if installation_total is not None
                else derive_total(installation_unit, quantity)
            )
        else:
            continue

        if all(
            value is None
            for value in (material_unit, material_total, installation_unit, installation_total)
        ):
            continue

        yield {
            "system": system,
            "excel_row": excel_row,
            "position": position,
            "name": name,
            "type_code_article": text_or_none(row[2]),
            "manufacturer": text_or_none(row[3]),
            "unit": unit,
            "quantity": quantity,
            "material_unit_kopeks": material_unit,
            "material_total_kopeks": material_total,
            "installation_unit_kopeks": installation_unit,
            "installation_total_kopeks": installation_total,
            "note": text_or_none(row[10]) or "лист: KP",
        }


def main():
    rows = list(iter_price_rows(SOURCE))
    source_stat = SOURCE.stat()
    source_hash = file_sha256(SOURCE)
    headers = {
        "sheet": "KP",
        "title": "Спецификация оборудования и материалов систем АПС и СОУЭ",
        "columns": [
            "поз.",
            "Наименование",
            "Тип / код / артикул",
            "Изготовитель",
            "Ед.",
            "Кол-во",
            "Материал: ед. расценка, руб. с НДС",
            "Материал: итого, руб. с НДС",
            "Монтаж: ед. расценка, руб. с НДС",
            "Монтаж: итого, руб. с НДС",
            "Примечание",
        ],
        "import_rule": "Rows 10-47 are imported as AПС/СОУЭ positions. Rows 50-54 are imported as separate misc expenses; totals are derived from unit price and quantity when the source only has a unit price.",
        "vat_rule": VAT_NOTE,
    }

    with sqlite3.connect(DB_PATH) as conn:
        conn.execute("PRAGMA foreign_keys = ON")
        previous = conn.execute(
            "SELECT id FROM price_snapshots WHERE source_file = ? OR source_sha256 = ?",
            (str(SOURCE), source_hash),
        ).fetchall()
        for (snapshot_id,) in previous:
            conn.execute("DELETE FROM price_snapshots WHERE id = ?", (snapshot_id,))

        cursor = conn.execute(
            """
            INSERT INTO price_snapshots (
                source_file,
                source_sha256,
                source_last_modified,
                price_date,
                imported_at,
                source_headers_json
            )
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (
                str(SOURCE),
                source_hash,
                dt.datetime.fromtimestamp(source_stat.st_mtime).isoformat(timespec="seconds"),
                PRICE_DATE,
                dt.datetime.now().isoformat(timespec="seconds"),
                json.dumps(headers, ensure_ascii=False),
            ),
        )
        snapshot_id = cursor.lastrowid

        for row in rows:
            cursor = conn.execute(
                """
                INSERT INTO price_lines (
                    snapshot_id,
                    system,
                    excel_row,
                    position,
                    name,
                    type_code_article,
                    manufacturer,
                    unit,
                    quantity,
                    note
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    snapshot_id,
                    row["system"],
                    row["excel_row"],
                    row["position"],
                    row["name"],
                    row["type_code_article"],
                    row["manufacturer"],
                    row["unit"],
                    row["quantity"],
                    row["note"],
                ),
            )
            line_id = cursor.lastrowid

            if row["material_unit_kopeks"] is not None or row["material_total_kopeks"] is not None:
                conn.execute(
                    """
                    INSERT INTO prices (
                        line_id,
                        price_kind,
                        unit_price_kopeks,
                        total_price_kopeks,
                        vat_note
                    )
                    VALUES (?, 'material', ?, ?, ?)
                    """,
                    (
                        line_id,
                        row["material_unit_kopeks"],
                        row["material_total_kopeks"],
                        VAT_NOTE,
                    ),
                )

            if row["installation_unit_kopeks"] is not None or row["installation_total_kopeks"] is not None:
                conn.execute(
                    """
                    INSERT INTO prices (
                        line_id,
                        price_kind,
                        unit_price_kopeks,
                        total_price_kopeks,
                        vat_note
                    )
                    VALUES (?, 'installation', ?, ?, ?)
                    """,
                    (
                        line_id,
                        row["installation_unit_kopeks"],
                        row["installation_total_kopeks"],
                        VAT_NOTE,
                    ),
                )

        result = {
            "snapshot_id": snapshot_id,
            "inserted_lines": conn.execute(
                "SELECT COUNT(*) FROM price_lines WHERE snapshot_id = ?", (snapshot_id,)
            ).fetchone()[0],
            "inserted_prices": [
                {
                    "price_kind": price_kind,
                    "records": records,
                    "total_rub": total_rub,
                }
                for price_kind, records, total_rub in conn.execute(
                    """
                    SELECT price_kind, COUNT(*), ROUND(SUM(total_price_kopeks) / 100.0, 2)
                    FROM prices
                    JOIN price_lines ON price_lines.id = prices.line_id
                    WHERE snapshot_id = ?
                    GROUP BY price_kind
                    ORDER BY price_kind
                    """,
                    (snapshot_id,),
                )
            ],
            "lines_by_system": [
                {"system": system, "lines": lines}
                for system, lines in conn.execute(
                    """
                    SELECT system, COUNT(*)
                    FROM price_lines
                    WHERE snapshot_id = ?
                    GROUP BY system
                    ORDER BY system
                    """,
                    (snapshot_id,),
                )
            ],
            "deleted_previous_snapshots": len(previous),
        }

    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
