from copy import copy
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs" / "Синтек 0806 с ценами с НДС.xlsx"
OUTPUT = ROOT / "outputs" / "Синтек 0806 с ценами с НДС - счет 45012 +30.xlsx"
SHEET_NAME = "АУП ПП"
REPORT_NAME = "Отчет сопоставления"
MARKUP = Decimal("1.30")


def money(value):
    return Decimal(str(value).replace(" ", "").replace(",", ".")).quantize(
        Decimal("0.01"), rounding=ROUND_HALF_UP
    )


def as_decimal(value):
    if value is None or value == "":
        return None
    return Decimal(str(value).replace(" ", "").replace(",", "."))


def is_blank(value):
    return value is None or value == ""


def append_note(existing, note):
    if is_blank(existing):
        return note
    text = str(existing)
    if note in text:
        return text
    return f"{text}; {note}"


def copy_number_style(src, dst):
    if src.has_style:
        dst._style = copy(src._style)
    dst.number_format = "# ##0.00"


invoice_items = {
    "XP7627": {
        "article": "XP7627",
        "name": 'ХОМУТ РЕЗЬБОВОЙ СЕДЁЛКА (30mm) 2 1/2"x3/4" - (76x27мм) Ду 65х20мм',
        "qty": Decimal("1218"),
        "price": money("446.01"),
    },
    "TMK76": {
        "article": "TMK76",
        "name": 'ТРОЙНИК ПОД МУФТУ 2 1/2" (76мм) Ду 65мм',
        "qty": Decimal("165"),
        "price": money("650.65"),
    },
    "MJ76_QTY165": {
        "article": "MJ76",
        "name": 'МУФТА ЖЕСТКАЯ 2 1/2" (76мм) Ду 65мм',
        "qty": Decimal("165"),
        "price": money("385.07"),
    },
    "PM10876": {
        "article": "PM10876",
        "name": 'ПЕРЕХОД ПОД МУФТУ 4"x2 1/2" (108х76мм) Ду 100х65мм',
        "qty": Decimal("165"),
        "price": money("403.28"),
    },
    "TPM219108": {
        "article": "TPM219108",
        "name": 'ТРОЙНИК ПЕРЕХОДНОЙ ПОД МУФТУ 8"x4" - (219x108мм) Ду 200x100мм',
        "qty": Decimal("165"),
        "price": money("5356.52"),
    },
    "TPM21976": {
        "article": "TPM21976",
        "name": 'ТРОЙНИК ПЕРЕХОДНОЙ ПОД МУФТУ 8"x2 1/2" - (219x76мм) Ду 200x65мм',
        "qty": Decimal("61"),
        "price": money("4189.14"),
    },
    "TMK219": {
        "article": "TMK219",
        "name": 'ТРОЙНИК ПОД МУФТУ 8" - (219мм) Ду 200мм',
        "qty": Decimal("51"),
        "price": money("5516.78"),
    },
    "KK90219": {
        "article": "KK90219",
        "name": 'КОЛЕНО 90° - 8" (219мм). ДУ 200мм',
        "qty": Decimal("10"),
        "price": money("3545.47"),
    },
    "KK9076": {
        "article": "KK9076",
        "name": 'КОЛЕНО 90° - 2 1/2" (76мм). ДУ 65мм',
        "qty": Decimal("63"),
        "price": money("402.06"),
    },
    "MJ219": {
        "article": "MJ219",
        "name": 'МУФТА ЖЕСТКАЯ 8" (219мм) Ду 200мм',
        "qty": Decimal("677"),
        "price": money("1620.82"),
    },
    "MJ108": {
        "article": "MJ108",
        "name": 'МУФТА ЖЕСТКАЯ 4" (108мм) Ду 100мм',
        "qty": Decimal("330"),
        "price": money("815.67"),
    },
    "MJ76_QTY974": {
        "article": "MJ76",
        "name": 'МУФТА ЖЕСТКАЯ 2 1/2" (76мм) Ду 65мм',
        "qty": Decimal("974"),
        "price": money("385.07"),
    },
}


row_map = {
    6: (["XP7627"], "аналог по назначению; количество в счете 1218 разделено на строки 6 и 7"),
    7: (["XP7627"], "аналог по назначению; количество в счете 1218 разделено на строки 6 и 7"),
    8: (["MJ76_QTY165", "TMK76", "PM10876"], "комплектный аналог по трем строкам счета с количеством 165"),
    9: (["TPM219108"], "аналог по типу фитинга, размеру и количеству"),
    10: (["TPM21976"], "аналог по типу фитинга, размеру и количеству"),
    11: (["TMK219"], "аналог по типу фитинга, размеру и количеству"),
    12: (["KK90219"], "аналог по типу фитинга, размеру и количеству"),
    13: (["KK9076"], "аналог по типу фитинга, размеру и количеству"),
    14: (["MJ219"], "аналог по типу фитинга, размеру и количеству"),
    15: (["MJ108"], "аналог по типу фитинга, размеру и количеству"),
    16: (["MJ76_QTY974"], "аналог по типу фитинга, размеру и количеству"),
}


def update_report(wb, updates):
    if REPORT_NAME not in wb.sheetnames:
        return
    report = wb[REPORT_NAME]
    by_row = {row: data for row, data in updates.items()}

    for r in range(2, report.max_row + 1):
        if report.cell(r, 2).value != SHEET_NAME:
            continue
        row_num = report.cell(r, 3).value
        if row_num not in by_row:
            continue

        data = by_row[row_num]
        items = data["items"]
        matched_name = " + ".join(item["name"] for item in items)
        matched_article = " + ".join(item["article"] for item in items)
        report.cell(r, 1).value = "заполнено по счету"
        report.cell(r, 8).value = float(data["unit"])
        report.cell(r, 9).value = None
        report.cell(r, 10).value = "Счет 45012 от 09.06.2026 / GROOVEX (+30%)"
        report.cell(r, 11).value = matched_name
        report.cell(r, 12).value = matched_article
        report.cell(r, 13).value = 0.9
        report.cell(r, 14).value = None
        report.cell(r, 15).value = "с НДС по счету; в документ поставлено +30%"
        report.cell(r, 16).value = data["comment"]


def main():
    wb = load_workbook(SOURCE)
    ws = wb[SHEET_NAME]

    updates = {}
    for row, (articles, match_comment) in row_map.items():
        items = [invoice_items[article] for article in articles]
        qty = as_decimal(ws.cell(row, 6).value)
        invoice_unit = sum((item["price"] for item in items), Decimal("0"))
        unit = money(invoice_unit * MARKUP)
        total = money(unit * qty) if qty is not None else None

        changed = []
        if is_blank(ws.cell(row, 7).value):
            ws.cell(row, 7).value = float(unit)
            copy_number_style(ws.cell(row - 1, 7), ws.cell(row, 7))
            changed.append("материал ед.")
        if total is not None and is_blank(ws.cell(row, 8).value):
            ws.cell(row, 8).value = float(total)
            copy_number_style(ws.cell(row - 1, 8), ws.cell(row, 8))
            changed.append("материал итого")

        note = (
            f"цена материала по счету 45012 от 09.06.2026 +30%; "
            f"арт. {' + '.join(item['article'] for item in items)}; {match_comment}"
        )
        if changed:
            ws.cell(row, 11).value = append_note(ws.cell(row, 11).value, note)
            updates[row] = {
                "unit": unit,
                "total": total,
                "items": items,
                "comment": match_comment,
                "changed": changed,
            }

    update_report(wb, updates)

    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
    except AttributeError:
        pass

    wb.save(OUTPUT)

    print(OUTPUT)
    for row, data in updates.items():
        print(row, "+".join(item["article"] for item in data["items"]), data["unit"], data["total"], data["changed"])


if __name__ == "__main__":
    main()
