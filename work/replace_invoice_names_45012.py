from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "outputs" / "Синтек 0806 с ценами с НДС - счет 45012 +30.xlsx"
OUTPUT = ROOT / "outputs" / "Синтек 0806 с ценами с НДС - счет 45012 +30 названия.xlsx"
SHEET_NAME = "АУП ПП"
REPORT_NAME = "Отчет сопоставления"


invoice_items = {
    "XP7627": {
        "article": "XP7627",
        "name": 'ХОМУТ РЕЗЬБОВОЙ СЕДЁЛКА (30mm) 2 1/2"x3/4" - (76x27мм) Ду 65х20мм. арт. XP7627 FM/UL GROOVEX',
    },
    "MJ76_QTY165": {
        "article": "MJ76",
        "name": 'МУФТА ЖЕСТКАЯ 2 1/2" (76мм) Ду 65мм. арт. MJ76 FM/UL GROOVEX',
    },
    "TMK76": {
        "article": "TMK76",
        "name": 'ТРОЙНИК ПОД МУФТУ 2 1/2" - (76мм) Ду 65мм. арт. TMK76 FM/UL GROOVEX',
    },
    "PM10876": {
        "article": "PM10876",
        "name": 'ПЕРЕХОД ПОД МУФТУ 4"x2 1/2" (108х76мм) Ду 100х65мм. арт. PM10876 FM/UL GROOVEX',
    },
    "TPM219108": {
        "article": "TPM219108",
        "name": 'ТРОЙНИК ПЕРЕХОДНОЙ ПОД МУФТУ 8"x4" - (219x108мм) Ду 200x100мм. арт. TPM219108 FM/UL GROOVEX',
    },
    "TPM21976": {
        "article": "TPM21976",
        "name": 'ТРОЙНИК ПЕРЕХОДНОЙ ПОД МУФТУ 8"x2 1/2" - (219x76мм) Ду 200x65мм. арт. TPM21976 FM/UL GROOVEX',
    },
    "TMK219": {
        "article": "TMK219",
        "name": 'ТРОЙНИК ПОД МУФТУ 8" - (219мм) Ду 200мм. арт. TMK219 FM/UL GROOVEX',
    },
    "KK90219": {
        "article": "KK90219",
        "name": 'КОЛЕНО 90° - 8" (219мм). ДУ 200мм. арт. KK90219 FM/UL GROOVEX',
    },
    "KK9076": {
        "article": "KK9076",
        "name": 'КОЛЕНО 90° - 2 1/2" (76мм). ДУ 65мм. арт. KK9076 FM/UL GROOVEX',
    },
    "MJ219": {
        "article": "MJ219",
        "name": 'МУФТА ЖЕСТКАЯ 8" (219мм) Ду 200мм. арт. MJ219 FM/UL GROOVEX',
    },
    "MJ108": {
        "article": "MJ108",
        "name": 'МУФТА ЖЕСТКАЯ 4" (108мм) Ду 100мм. арт. MJ108 FM/UL GROOVEX',
    },
    "MJ76_QTY974": {
        "article": "MJ76",
        "name": 'МУФТА ЖЕСТКАЯ 2 1/2" (76мм) Ду 65мм. арт. MJ76 FM/UL GROOVEX',
    },
}


row_map = {
    6: ["XP7627"],
    7: ["XP7627"],
    8: ["MJ76_QTY165", "TMK76", "PM10876"],
    9: ["TPM219108"],
    10: ["TPM21976"],
    11: ["TMK219"],
    12: ["KK90219"],
    13: ["KK9076"],
    14: ["MJ219"],
    15: ["MJ108"],
    16: ["MJ76_QTY974"],
}


def combined_name(keys):
    items = [invoice_items[key] for key in keys]
    if len(items) == 1:
        return items[0]["name"]
    return "Комплект по счету 45012: " + " + ".join(item["name"] for item in items)


def combined_article(keys):
    return " + ".join(invoice_items[key]["article"] for key in keys)


def append_note(value, note):
    if value is None or value == "":
        return note
    text = str(value)
    if note in text:
        return text
    return f"{text}; {note}"


def sync_report(wb):
    if REPORT_NAME not in wb.sheetnames:
        return
    ws = wb[REPORT_NAME]
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, 2).value != SHEET_NAME:
            continue
        target_row = ws.cell(row, 3).value
        if target_row not in row_map:
            continue
        keys = row_map[target_row]
        ws.cell(row, 5).value = combined_name(keys)
        ws.cell(row, 6).value = combined_article(keys)
        ws.cell(row, 11).value = combined_name(keys)
        ws.cell(row, 12).value = combined_article(keys)
        ws.cell(row, 16).value = append_note(
            ws.cell(row, 16).value,
            "наименование и артикул заменены на данные из счета",
        )


def main():
    wb = load_workbook(INPUT)
    ws = wb[SHEET_NAME]

    for row, keys in row_map.items():
        ws.cell(row, 2).value = combined_name(keys)
        ws.cell(row, 3).value = combined_article(keys)
        ws.cell(row, 4).value = "GROOVEX"
        ws.cell(row, 11).value = append_note(
            ws.cell(row, 11).value,
            "наименование и артикул заменены на данные из счета",
        )

    sync_report(wb)
    wb.save(OUTPUT)

    print(OUTPUT)
    for row in row_map:
        print(row, ws.cell(row, 2).value, ws.cell(row, 3).value, ws.cell(row, 4).value)


if __name__ == "__main__":
    main()
